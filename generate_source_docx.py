from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── Page Margins ──────────────────────────────────────────────────────────────
for sec in doc.sections:
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.2); sec.right_margin = Inches(1)

NAVY = RGBColor(0x1a, 0x23, 0x7e)
GOLD = RGBColor(0xf9, 0xa8, 0x25)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x37, 0x47, 0x4F)
CODEBLUE = RGBColor(0x00, 0x3a, 0x6e)

def heading1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.size = Pt(16); run.font.color.rgb = NAVY; run.bold = True
    return p

def heading2(text, color=GOLD):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(13); run.font.color.rgb = color; run.bold = True
    return p

def body(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11); run.font.color.rgb = GRAY
    p.paragraph_format.space_after = Pt(4)
    return p

def code_block(filepath):
    if not os.path.exists(filepath):
        body(f"[File not found: {filepath}]"); return
    with open(filepath, encoding='utf-8') as f:
        lines = f.readlines()
    p_header = doc.add_paragraph()
    r = p_header.add_run(f"  📄 {os.path.basename(filepath)}  ({len(lines)} lines)")
    r.font.size = Pt(9); r.font.color.rgb = CODEBLUE; r.bold = True
    p_header.paragraph_format.space_after = Pt(0)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.3)
        r = p.add_run(line.rstrip())
        r.font.name = 'Courier New'; r.font.size = Pt(8.5); r.font.color.rgb = BLACK
    doc.add_paragraph()

BASE = r"C:\Users\Rutuja\OneDrive\Desktop\MINI PROJECT\SUREKHA"

# ── TITLE PAGE ─────────────────────────────────────────────────────────────────
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("SUREKHA")
r.font.size = Pt(36); r.font.color.rgb = NAVY; r.bold = True
doc.add_paragraph()

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = t2.add_run("Smart Urban Residential Estate &\nKnowledge-based Housing Accounting")
r2.font.size = Pt(14); r2.font.color.rgb = GOLD; r2.bold = True
doc.add_paragraph()

t3 = doc.add_paragraph()
t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = t3.add_run("COMPLETE SOURCE CODE DOCUMENT")
r3.font.size = Pt(13); r3.font.color.rgb = GRAY; r3.bold = True
doc.add_paragraph()

t4 = doc.add_paragraph()
t4.alignment = WD_ALIGN_PARAGRAPH.CENTER
r4 = t4.add_run("TYCS – Semester V  |  Mini Project  |  Python Flask + SQLite")
r4.font.size = Pt(11); r4.font.color.rgb = GRAY
doc.add_page_break()

# ── MODULE OVERVIEW ────────────────────────────────────────────────────────────
heading1("Project Module Overview")
modules = [
    ("Module 1", "Authentication & Society Setup",
     "Firm Registration, Login, Email OTP Password Reset, Society Creation & Selection. Files: auth.py, society.py"),
    ("Module 2", "Members & Statutory Registers",
     "Member Add/Delete, Form I (Membership), Form J (Share Transfer), Share Certificate Print, Email/SMS Broadcast. Files: members.py, registers.py"),
    ("Module 3", "Billing & Financial Accounting",
     "19-Item Receipt Entry, Payment Voucher, Journal Entry, Ledger, Cash/Bank Book, Maintenance Bill Generation & Print. Files: entries.py, billing.py"),
    ("Module 4", "Final Accounts & Reporting",
     "Income & Expenditure Account, Balance Sheet, Outstanding Dues Report, Member Summary. Files: accounts.py, reports.py"),
    ("Module 5", "Communication & Digital Meetings",
     "AGM/SGM Meeting Schedule, Jitsi Video Conference Link, Agenda Points, Notice Board. Files: meetings.py"),
]
for code, name, desc in modules:
    p = doc.add_paragraph()
    r0 = p.add_run(f"  {code}: ")
    r0.font.size = Pt(11); r0.font.color.rgb = GOLD; r0.bold = True
    r = p.add_run(name)
    r.font.size = Pt(11); r.font.color.rgb = NAVY; r.bold = True
    body(f"     {desc}")
doc.add_page_break()

# ── SOURCE FILES ───────────────────────────────────────────────────────────────
files = [
    ("app.py", "Main Application Entry Point",
     "Registers all 10 Flask Blueprints and initializes the SQLite database on startup."),
    ("database.py", "Database Schema & Connection",
     "Defines all 13 SQLite tables and provides get_db() / init_db() helper functions."),
    ("routes/auth.py", "Module 1 – Authentication",
     "Handles Login, Firm Registration, Email OTP Forgot Password, Dashboard summary."),
    ("routes/society.py", "Module 1 – Society Management",
     "Add, Select, and Delete societies. Sets society_id in Flask session."),
    ("routes/members.py", "Module 2 – Members & Registers",
     "Add/Delete Members, Form I/J entries, Broadcast Email+SMS notifications."),
    ("routes/registers.py", "Module 2 – Statutory Registers Print",
     "Print Form I, Form J, and Individual Share Certificate as HTML pages."),
    ("routes/entries.py", "Module 3 – Financial Entries",
     "19-item Receipts, Payments, Journal, Ledger, and Cash/Bank Book CRUD operations."),
    ("routes/billing.py", "Module 3 – Billing & num_words()",
     "Configure charges, Auto/Manual bill generation, Print bill with amount-in-words helper."),
    ("routes/accounts.py", "Module 4 – Final Accounts",
     "Income & Expenditure statement with running surplus/deficit calculation."),
    ("routes/reports.py", "Module 4 – Reports",
     "Outstanding dues report and complete member summary listing."),
    ("routes/meetings.py", "Module 5 – Meetings & Video",
     "Schedule AGM/SGM meetings, auto-generate Jitsi video links, manage agenda/notices."),
]

for fname, title, desc in files:
    heading1(f"{'=' * 3} {title}")
    body(desc)
    code_block(os.path.join(BASE, fname.replace('/', os.sep)))
    doc.add_page_break()

# ── SAVE ───────────────────────────────────────────────────────────────────────
OUT1 = r"C:\Users\Rutuja\OneDrive\Desktop\SUREKHA_SOURCE_CODE.docx"
OUT2 = os.path.join(BASE, "SUREKHA_SOURCE_CODE.docx")
doc.save(OUT1); doc.save(OUT2)
print(f"DOCX saved!\n1. {OUT1}\n2. {OUT2}")
print(f"   Size: {os.path.getsize(OUT1):,} bytes")
